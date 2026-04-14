"""
E2E tests for GET /api/v1/reports/content-export.

Uses SQLite in-memory with create_all so no real Postgres is needed.
The service.build_content_export call is mocked to isolate HTTP-layer
behaviour (auth, date validation, response headers) from Excel generation.
Integration-style data tests use the repository directly with real fixtures.
"""

import uuid
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import AsyncMock, patch

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.auth.models import Organization, User
from app.catalog.models import Brand, Platform, SKU, SKUPlatform
from app.core.database import Base
from app.core.deps import get_db
from app.core.security import create_access_token, hash_password
from app.main import app
from app.reports.models import ContentScoreRead

TEST_DATABASE_URL = "sqlite+aiosqlite:///:memory:"

ORG_A_ID = uuid.uuid4()
ORG_B_ID = uuid.uuid4()

_FAKE_XLSX = b"PK\x03\x04"  # minimal placeholder for mocked responses


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest_asyncio.fixture(scope="session")
async def engine():
    e = create_async_engine(TEST_DATABASE_URL, echo=False)
    async with e.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    yield e
    await e.dispose()


@pytest_asyncio.fixture
async def db_session(engine):
    async with engine.connect() as conn:
        await conn.begin()
        factory = async_sessionmaker(bind=conn, expire_on_commit=False, class_=AsyncSession)
        async with factory() as session:
            yield session
        await conn.rollback()


@pytest_asyncio.fixture
async def client(db_session):
    async def _override():
        yield db_session

    app.dependency_overrides[get_db] = _override
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c
    app.dependency_overrides.clear()


@pytest_asyncio.fixture
async def org_a(db_session):
    org = Organization(id=ORG_A_ID, name="Org A", slug="org-a", plan="pro")
    db_session.add(org)
    await db_session.flush()
    return org


@pytest_asyncio.fixture
async def org_b(db_session):
    org = Organization(id=ORG_B_ID, name="Org B", slug="org-b", plan="basic")
    db_session.add(org)
    await db_session.flush()
    return org


@pytest_asyncio.fixture
async def manager_user(db_session, org_a):
    user = User(
        id=uuid.uuid4(),
        org_id=org_a.id,
        email="manager@org-a.com",
        password_hash=hash_password("pass"),
        role="manager",
    )
    db_session.add(user)
    await db_session.flush()
    return user


@pytest_asyncio.fixture
async def viewer_user(db_session, org_a):
    user = User(
        id=uuid.uuid4(),
        org_id=org_a.id,
        email="viewer@org-a.com",
        password_hash=hash_password("pass"),
        role="viewer",
    )
    db_session.add(user)
    await db_session.flush()
    return user


@pytest_asyncio.fixture
async def org_b_user(db_session, org_b):
    user = User(
        id=uuid.uuid4(),
        org_id=org_b.id,
        email="user@org-b.com",
        password_hash=hash_password("pass"),
        role="manager",
    )
    db_session.add(user)
    await db_session.flush()
    return user


def _auth_headers(user: User) -> dict:
    token = create_access_token(user.id, user.org_id, user.role)
    return {"Authorization": f"Bearer {token}"}


# ---------------------------------------------------------------------------
# Auth & parameter validation tests (service mocked)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_export_requires_authentication(client):
    resp = await client.get(
        "/api/v1/reports/content-export",
        params={"date_from": "2026-01-01", "date_to": "2026-01-31"},
    )
    assert resp.status_code == 401


@pytest.mark.asyncio
async def test_export_missing_date_params_returns_422(client, manager_user):
    resp = await client.get(
        "/api/v1/reports/content-export",
        headers=_auth_headers(manager_user),
    )
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_export_date_to_before_date_from_returns_400(client, manager_user):
    resp = await client.get(
        "/api/v1/reports/content-export",
        params={"date_from": "2026-03-01", "date_to": "2026-01-01"},
        headers=_auth_headers(manager_user),
    )
    assert resp.status_code == 400
    assert "INVALID_DATE_RANGE" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_export_date_range_exceeds_limit_returns_400(client, manager_user):
    resp = await client.get(
        "/api/v1/reports/content-export",
        params={"date_from": "2024-01-01", "date_to": "2026-01-01"},
        headers=_auth_headers(manager_user),
    )
    assert resp.status_code == 400
    assert "INVALID_DATE_RANGE" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_export_returns_xlsx_content_type(client, manager_user):
    fake_bytes = _make_real_xlsx()
    with patch(
        "app.reports.router.service.build_content_export",
        new_callable=AsyncMock,
        return_value=fake_bytes,
    ):
        resp = await client.get(
            "/api/v1/reports/content-export",
            params={"date_from": "2026-01-01", "date_to": "2026-01-31"},
            headers=_auth_headers(manager_user),
        )
    assert resp.status_code == 200
    assert (
        resp.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "content-disposition" in resp.headers
    assert "content_scores_2026-01-01_2026-01-31.xlsx" in resp.headers["content-disposition"]


@pytest.mark.asyncio
async def test_export_viewer_role_allowed(client, viewer_user):
    fake_bytes = _make_real_xlsx()
    with patch(
        "app.reports.router.service.build_content_export",
        new_callable=AsyncMock,
        return_value=fake_bytes,
    ):
        resp = await client.get(
            "/api/v1/reports/content-export",
            params={"date_from": "2026-01-01", "date_to": "2026-01-31"},
            headers=_auth_headers(viewer_user),
        )
    assert resp.status_code == 200


# ---------------------------------------------------------------------------
# Integration test: repository returns only org-scoped data
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_content_export_cross_tenant_isolation(db_session, org_a, org_b):
    """
    Scores belonging to org_b must NOT appear in an export for org_a.
    """
    from app.reports.repository import get_content_scores_for_export

    platform = Platform(id=uuid.uuid4(), name="WB", type="marketplace")
    db_session.add(platform)

    brand_a = Brand(id=uuid.uuid4(), org_id=org_a.id, name="Brand A")
    brand_b = Brand(id=uuid.uuid4(), org_id=org_b.id, name="Brand B")
    db_session.add_all([brand_a, brand_b])

    sku_a = SKU(
        id=uuid.uuid4(),
        org_id=org_a.id,
        brand_id=brand_a.id,
        article="A-001",
        name="SKU A",
    )
    sku_b = SKU(
        id=uuid.uuid4(),
        org_id=org_b.id,
        brand_id=brand_b.id,
        article="B-001",
        name="SKU B",
    )
    db_session.add_all([sku_a, sku_b])

    sp_a = SKUPlatform(id=uuid.uuid4(), sku_id=sku_a.id, platform_id=platform.id)
    sp_b = SKUPlatform(id=uuid.uuid4(), sku_id=sku_b.id, platform_id=platform.id)
    db_session.add_all([sp_a, sp_b])

    from datetime import datetime, timezone

    score_a = ContentScoreRead(
        id=uuid.uuid4(),
        sku_platform_id=sp_a.id,
        scored_at=date(2026, 1, 15),
        content_total=Decimal("85.00"),
        created_at=datetime.now(tz=timezone.utc),
    )
    score_b = ContentScoreRead(
        id=uuid.uuid4(),
        sku_platform_id=sp_b.id,
        scored_at=date(2026, 1, 15),
        content_total=Decimal("72.00"),
        created_at=datetime.now(tz=timezone.utc),
    )
    db_session.add_all([score_a, score_b])
    await db_session.flush()

    rows = await get_content_scores_for_export(
        db=db_session,
        org_id=org_a.id,
        date_from=date(2026, 1, 1),
        date_to=date(2026, 1, 31),
    )

    assert len(rows) == 1
    assert rows[0].sku_article == "A-001"
    assert rows[0].brand_name == "Brand A"


# ---------------------------------------------------------------------------
# Unit test: Excel workbook structure
# ---------------------------------------------------------------------------


def _make_real_xlsx() -> bytes:
    """Build a minimal real xlsx for content-type and header tests."""
    from openpyxl import Workbook as WB

    wb = WB()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def test_build_workbook_headers():
    from app.reports.service import _build_workbook

    rows = []
    wb = _build_workbook(rows, date_from=date(2026, 1, 1), date_to=date(2026, 1, 31))
    ws = wb.active

    # Row 1 is the period sub-header
    assert "2026-01-01" in (ws["A1"].value or "")

    # Row 2 has column headers
    headers = [ws.cell(row=2, column=c).value for c in range(1, 10)]
    assert "Бренд" in headers
    assert "Артикул" in headers
    assert "Итого контент" in headers


def test_build_workbook_color_coding():
    from app.reports.service import ContentScoreRow, _build_workbook

    rows = [
        ContentScoreRow(
            brand_name="B",
            sku_article="001",
            sku_name="Name",
            platform_name="WB",
            scored_at=date(2026, 1, 1),
            image_score=Decimal("90"),
            description_score=Decimal("85"),
            composition_score=Decimal("80"),
            content_total=Decimal("85"),  # green
        ),
        ContentScoreRow(
            brand_name="B",
            sku_article="002",
            sku_name="Name2",
            platform_name="WB",
            scored_at=date(2026, 1, 1),
            image_score=Decimal("60"),
            description_score=Decimal("55"),
            composition_score=Decimal("50"),
            content_total=Decimal("55"),  # yellow
        ),
        ContentScoreRow(
            brand_name="B",
            sku_article="003",
            sku_name="Name3",
            platform_name="WB",
            scored_at=date(2026, 1, 1),
            image_score=Decimal("30"),
            description_score=Decimal("25"),
            composition_score=Decimal("20"),
            content_total=Decimal("25"),  # red
        ),
    ]
    wb = _build_workbook(rows, date_from=date(2026, 1, 1), date_to=date(2026, 1, 31))
    ws = wb.active

    # content_total column is column 9
    green_fill = ws.cell(row=3, column=9).fill.fgColor.rgb
    yellow_fill = ws.cell(row=4, column=9).fill.fgColor.rgb
    red_fill = ws.cell(row=5, column=9).fill.fgColor.rgb

    assert "C6EFCE" in green_fill
    assert "FFEB9C" in yellow_fill
    assert "FFC7CE" in red_fill


def test_build_workbook_null_scores_no_fill():
    from app.reports.service import ContentScoreRow, _build_workbook

    from datetime import datetime, timezone

    rows = [
        ContentScoreRow(
            brand_name="B",
            sku_article="001",
            sku_name="Name",
            platform_name="WB",
            scored_at=date(2026, 1, 1),
            image_score=None,
            description_score=None,
            composition_score=None,
            content_total=None,
        ),
    ]
    wb = _build_workbook(rows, date_from=date(2026, 1, 1), date_to=date(2026, 1, 31))
    ws = wb.active
    # No fill should be applied (default fill type is "none")
    assert ws.cell(row=3, column=9).fill.fill_type in (None, "none", "patternType")
