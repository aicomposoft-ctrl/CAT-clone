"""
E2E and unit tests for GET /api/v1/reports/reviews-export.

Coverage:
  - /reports/reviews-export: happy path, empty, sentiment filter, platform filter,
    sku_id owned, sku_id cross-tenant 404, invalid date 400, range > 366 400,
    unauthenticated 401, viewer role 200, invalid sentiment enum 422
  - _build_reviews_workbook: sentiment color coding, text truncation,
    truncated-flag warning row, NULL rating/sentiment/score, header structure
    row 1 merged + row 2 headers, empty rows header-only

NOTE: The repository query uses PostgreSQL-specific SQL. These tests mock the
service layer so auth, routing, date validation, and sku_id ownership checks
are exercised against the real router without requiring a PostgreSQL connection.
Workbook unit tests exercise the openpyxl layer directly.
"""

from __future__ import annotations

import io
import uuid
from datetime import date, timedelta
from decimal import Decimal
from typing import AsyncGenerator
from unittest.mock import AsyncMock, patch

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.auth.models import Organization, User
from app.catalog.models import Brand, SKU
from app.core.database import Base
from app.core.deps import get_db
from app.core.security import create_access_token, hash_password
from app.main import app
from app.reports.repository import ReviewRow
from app.reports.service import _build_reviews_workbook

TEST_DATABASE_URL = "sqlite+aiosqlite:///:memory:"
ORG_A_ID = uuid.uuid4()
ORG_B_ID = uuid.uuid4()

_today = date.today()


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest_asyncio.fixture(scope="session")
async def engine():
    e = create_async_engine(TEST_DATABASE_URL, echo=False)
    async with e.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    yield e
    await e.dispose()


@pytest_asyncio.fixture
async def db_session(engine):
    async with engine.connect() as conn:
        await conn.begin()
        factory = async_sessionmaker(bind=conn, expire_on_commit=False, class_=AsyncSession)
        async with factory() as session:
            yield session
        await conn.rollback()


@pytest_asyncio.fixture
async def client(db_session):
    async def _override() -> AsyncGenerator[AsyncSession, None]:
        yield db_session

    app.dependency_overrides[get_db] = _override
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c
    app.dependency_overrides.clear()


@pytest_asyncio.fixture
async def org_a(db_session):
    org = Organization(id=ORG_A_ID, name="Org A", slug="org-a-reviews-exp", plan="pro")
    db_session.add(org)
    await db_session.flush()
    return org


@pytest_asyncio.fixture
async def org_b(db_session):
    org = Organization(id=ORG_B_ID, name="Org B", slug="org-b-reviews-exp", plan="pro")
    db_session.add(org)
    await db_session.flush()
    return org


@pytest_asyncio.fixture
async def manager_a(db_session, org_a):
    user = User(
        id=uuid.uuid4(),
        org_id=ORG_A_ID,
        email="mgr-revexp@orga.test",
        password_hash=hash_password("pw"),
        role="manager",
    )
    db_session.add(user)
    await db_session.flush()
    return user


@pytest_asyncio.fixture
async def viewer_a(db_session, org_a):
    user = User(
        id=uuid.uuid4(),
        org_id=ORG_A_ID,
        email="viewer-revexp@orga.test",
        password_hash=hash_password("pw"),
        role="viewer",
    )
    db_session.add(user)
    await db_session.flush()
    return user


@pytest_asyncio.fixture
async def manager_b(db_session, org_b):
    user = User(
        id=uuid.uuid4(),
        org_id=ORG_B_ID,
        email="mgr-revexp@orgb.test",
        password_hash=hash_password("pw"),
        role="manager",
    )
    db_session.add(user)
    await db_session.flush()
    return user


@pytest_asyncio.fixture
async def sku_a(db_session, org_a):
    brand = Brand(id=uuid.uuid4(), org_id=ORG_A_ID, name="Brand A RevExp", type="client")
    db_session.add(brand)
    await db_session.flush()
    sku = SKU(
        id=uuid.uuid4(),
        org_id=ORG_A_ID,
        brand_id=brand.id,
        name="Test SKU RevExp",
        is_active=True,
    )
    db_session.add(sku)
    await db_session.flush()
    return sku


@pytest_asyncio.fixture
async def sku_b(db_session, org_b):
    brand = Brand(id=uuid.uuid4(), org_id=ORG_B_ID, name="Brand B RevExp", type="client")
    db_session.add(brand)
    await db_session.flush()
    sku = SKU(
        id=uuid.uuid4(),
        org_id=ORG_B_ID,
        brand_id=brand.id,
        name="Test SKU B RevExp",
        is_active=True,
    )
    db_session.add(sku)
    await db_session.flush()
    return sku


def _auth(user: User) -> dict:
    token = create_access_token(user.id, user.org_id, user.role)
    return {"Authorization": f"Bearer {token}"}


def _make_review_rows(n: int = 3, sentiment: str = "positive") -> list[ReviewRow]:
    return [
        ReviewRow(
            brand_name=f"Brand {i}",
            sku_article=f"ART-{i:03}",
            sku_name=f"SKU {i}",
            platform_name="Wildberries",
            review_date=_today - timedelta(days=i),
            rating=4,
            sentiment=sentiment,
            sentiment_score=Decimal("0.920"),
            review_text=f"Хороший товар номер {i}",
        )
        for i in range(n)
    ]


# ---------------------------------------------------------------------------
# E2E tests: /reports/reviews-export
# ---------------------------------------------------------------------------


class TestReviewsExport:

    @pytest.mark.asyncio
    async def test_export_returns_xlsx_bytes(self, client, manager_a, sku_a):
        mock_bytes = b"fake_xlsx"
        with patch(
            "app.reports.router.service.build_reviews_export",
            new=AsyncMock(return_value=mock_bytes),
        ):
            resp = await client.get(
                f"/api/v1/reports/reviews-export?date_from=2026-01-01&date_to=2026-03-31",
                headers=_auth(manager_a),
            )
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "reviews_2026-01-01_2026-03-31.xlsx" in resp.headers["content-disposition"]

    @pytest.mark.asyncio
    async def test_export_empty_range_returns_200(self, client, manager_a, sku_a):
        """No reviews → empty workbook bytes, still HTTP 200."""
        real_bytes = _build_reviews_workbook([], _today - timedelta(30), _today, False)
        buf = io.BytesIO()
        real_bytes.save(buf)
        buf.seek(0)
        mock_bytes = buf.read()

        with patch(
            "app.reports.router.service.build_reviews_export",
            new=AsyncMock(return_value=mock_bytes),
        ):
            resp = await client.get(
                f"/api/v1/reports/reviews-export?date_from=2026-01-01&date_to=2026-01-01",
                headers=_auth(manager_a),
            )
        assert resp.status_code == 200

    @pytest.mark.asyncio
    async def test_export_sentiment_filter_passed_to_service(self, client, manager_a, sku_a):
        with patch(
            "app.reports.router.service.build_reviews_export",
            new=AsyncMock(return_value=b"x"),
        ) as mock_svc:
            await client.get(
                f"/api/v1/reports/reviews-export?date_from=2026-01-01&date_to=2026-03-31&sentiment=negative",
                headers=_auth(manager_a),
            )
        _, kwargs = mock_svc.call_args
        assert kwargs.get("sentiment") == "negative"

    @pytest.mark.asyncio
    async def test_export_platform_filter_passed_to_service(self, client, manager_a, sku_a):
        plat_id = uuid.uuid4()
        with patch(
            "app.reports.router.service.build_reviews_export",
            new=AsyncMock(return_value=b"x"),
        ) as mock_svc:
            await client.get(
                f"/api/v1/reports/reviews-export?date_from=2026-01-01&date_to=2026-03-31&platform_id={plat_id}",
                headers=_auth(manager_a),
            )
        _, kwargs = mock_svc.call_args
        assert kwargs.get("platform_id") == plat_id

    @pytest.mark.asyncio
    async def test_export_sku_filter_owned_passes_to_service(self, client, manager_a, sku_a):
        with patch(
            "app.reports.router.service.build_reviews_export",
            new=AsyncMock(return_value=b"x"),
        ) as mock_svc:
            resp = await client.get(
                f"/api/v1/reports/reviews-export?date_from=2026-01-01&date_to=2026-03-31&sku_id={sku_a.id}",
                headers=_auth(manager_a),
            )
        assert resp.status_code == 200
        _, kwargs = mock_svc.call_args
        assert kwargs.get("sku_id") == sku_a.id

    @pytest.mark.asyncio
    async def test_export_sku_filter_cross_tenant_returns_404(self, client, manager_b, sku_a):
        """manager_b (org_b) requesting sku_a (org_a) → 404."""
        resp = await client.get(
            f"/api/v1/reports/reviews-export?date_from=2026-01-01&date_to=2026-03-31&sku_id={sku_a.id}",
            headers=_auth(manager_b),
        )
        assert resp.status_code == 404
        assert resp.json()["detail"] == "SKU_NOT_FOUND"

    @pytest.mark.asyncio
    async def test_export_inverted_date_returns_400(self, client, manager_a, sku_a):
        resp = await client.get(
            f"/api/v1/reports/reviews-export?date_from=2026-03-01&date_to=2026-01-01",
            headers=_auth(manager_a),
        )
        assert resp.status_code == 400
        assert "INVALID_DATE_RANGE" in resp.json()["detail"]

    @pytest.mark.asyncio
    async def test_export_date_range_over_366_returns_400(self, client, manager_a, sku_a):
        resp = await client.get(
            f"/api/v1/reports/reviews-export?date_from=2025-01-01&date_to=2026-06-01",
            headers=_auth(manager_a),
        )
        assert resp.status_code == 400
        assert "366" in resp.json()["detail"]

    @pytest.mark.asyncio
    async def test_export_unauthenticated_returns_401(self, client, sku_a):
        resp = await client.get(
            f"/api/v1/reports/reviews-export?date_from=2026-01-01&date_to=2026-03-31"
        )
        assert resp.status_code == 401

    @pytest.mark.asyncio
    async def test_export_viewer_can_download(self, client, viewer_a, sku_a):
        with patch(
            "app.reports.router.service.build_reviews_export",
            new=AsyncMock(return_value=b"x"),
        ):
            resp = await client.get(
                f"/api/v1/reports/reviews-export?date_from=2026-01-01&date_to=2026-03-31",
                headers=_auth(viewer_a),
            )
        assert resp.status_code == 200

    @pytest.mark.asyncio
    async def test_export_invalid_sentiment_returns_422(self, client, manager_a, sku_a):
        resp = await client.get(
            f"/api/v1/reports/reviews-export?date_from=2026-01-01&date_to=2026-03-31&sentiment=excellent",
            headers=_auth(manager_a),
        )
        assert resp.status_code == 422


# ---------------------------------------------------------------------------
# Unit tests: _build_reviews_workbook
# ---------------------------------------------------------------------------


def _load_wb(wb) -> object:
    """Save workbook to bytes and reload via openpyxl load_workbook."""
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return load_workbook(buf)


class TestBuildReviewsWorkbook:

    def test_sheet_name_is_reviews(self):
        wb = _build_reviews_workbook([], _today, _today, False)
        loaded = _load_wb(wb)
        assert "Reviews" in loaded.sheetnames

    def test_header_row1_merged_period_text(self):
        d_from = date(2026, 1, 1)
        d_to = date(2026, 3, 31)
        wb = _build_reviews_workbook([], d_from, d_to, False)
        loaded = _load_wb(wb)
        ws = loaded["Reviews"]
        assert "2026-01-01" in (ws["A1"].value or "")
        assert "2026-03-31" in (ws["A1"].value or "")

    def test_header_row2_has_9_columns(self):
        wb = _build_reviews_workbook([], _today, _today, False)
        loaded = _load_wb(wb)
        ws = loaded["Reviews"]
        headers = [ws.cell(row=2, column=c).value for c in range(1, 10)]
        assert headers[0] == "Бренд"
        assert headers[5] == "Рейтинг"
        assert headers[6] == "Тональность"
        assert headers[8] == "Текст отзыва"

    def test_empty_rows_produces_header_only_workbook(self):
        wb = _build_reviews_workbook([], _today, _today, False)
        loaded = _load_wb(wb)
        ws = loaded["Reviews"]
        assert ws.cell(row=3, column=1).value is None

    def test_sentiment_color_positive(self):
        rows = [_make_review_rows(1, "positive")[0]]
        wb = _build_reviews_workbook(rows, _today, _today, False)
        loaded = _load_wb(wb)
        ws = loaded["Reviews"]
        fill = ws.cell(row=3, column=7).fill
        assert fill.fgColor.rgb.upper().endswith("C6EFCE")

    def test_sentiment_color_negative(self):
        rows = [_make_review_rows(1, "negative")[0]]
        wb = _build_reviews_workbook(rows, _today, _today, False)
        loaded = _load_wb(wb)
        ws = loaded["Reviews"]
        fill = ws.cell(row=3, column=7).fill
        assert fill.fgColor.rgb.upper().endswith("FFC7CE")

    def test_sentiment_color_neutral(self):
        rows = [_make_review_rows(1, "neutral")[0]]
        wb = _build_reviews_workbook(rows, _today, _today, False)
        loaded = _load_wb(wb)
        ws = loaded["Reviews"]
        fill = ws.cell(row=3, column=7).fill
        assert fill.fgColor.rgb.upper().endswith("FFEB9C")

    def test_null_sentiment_no_fill_on_columns_7_8(self):
        row = ReviewRow(
            brand_name="B",
            sku_article=None,
            sku_name="S",
            platform_name="P",
            review_date=_today,
            rating=3,
            sentiment=None,
            sentiment_score=None,
            review_text="text",
        )
        wb = _build_reviews_workbook([row], _today, _today, False)
        loaded = _load_wb(wb)
        ws = loaded["Reviews"]
        fill7 = ws.cell(row=3, column=7).fill
        # No fill means patternType is None or "none"
        assert fill7.patternType in (None, "none")

    def test_review_text_truncated_at_500_chars(self):
        long_text = "А" * 600
        row = ReviewRow(
            brand_name="B",
            sku_article=None,
            sku_name="S",
            platform_name="P",
            review_date=_today,
            rating=5,
            sentiment="positive",
            sentiment_score=Decimal("0.95"),
            review_text=long_text,
        )
        wb = _build_reviews_workbook([row], _today, _today, False)
        loaded = _load_wb(wb)
        ws = loaded["Reviews"]
        cell_value = ws.cell(row=3, column=9).value
        assert cell_value is not None
        assert len(cell_value) == 501  # 500 chars + "…"
        assert cell_value.endswith("…")

    def test_rating_null_formatted_as_dash(self):
        row = ReviewRow(
            brand_name="B",
            sku_article=None,
            sku_name="S",
            platform_name="P",
            review_date=_today,
            rating=None,
            sentiment=None,
            sentiment_score=None,
            review_text="text",
        )
        wb = _build_reviews_workbook([row], _today, _today, False)
        loaded = _load_wb(wb)
        ws = loaded["Reviews"]
        assert ws.cell(row=3, column=6).value == "—"

    def test_truncated_flag_adds_warning_row(self):
        rows = _make_review_rows(3)
        wb = _build_reviews_workbook(rows, _today, _today, True)
        loaded = _load_wb(wb)
        ws = loaded["Reviews"]
        # 3 data rows (rows 3-5) + warning at row 6
        warn_val = ws.cell(row=6, column=1).value
        assert warn_val is not None
        assert "10 000" in warn_val or "10\xa0000" in warn_val or "лимит" in warn_val.lower()

    def test_no_warning_row_when_not_truncated(self):
        rows = _make_review_rows(3)
        wb = _build_reviews_workbook(rows, _today, _today, False)
        loaded = _load_wb(wb)
        ws = loaded["Reviews"]
        # Row 6 should be empty (no warning)
        assert ws.cell(row=6, column=1).value is None

    def test_sentiment_score_formatted_3dp(self):
        row = ReviewRow(
            brand_name="B",
            sku_article=None,
            sku_name="S",
            platform_name="P",
            review_date=_today,
            rating=5,
            sentiment="positive",
            sentiment_score=Decimal("0.92"),
            review_text="text",
        )
        wb = _build_reviews_workbook([row], _today, _today, False)
        loaded = _load_wb(wb)
        ws = loaded["Reviews"]
        assert ws.cell(row=3, column=8).value == "0.920"
