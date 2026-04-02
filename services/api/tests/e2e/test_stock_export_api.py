"""
E2E tests for GET /api/v1/reports/stock-export.

HTTP-layer tests mock service.build_stock_export to isolate auth, date
validation, and response header behaviour.  Workbook-structure tests are
pure unit tests that call _build_stock_workbook directly.
"""

import uuid
from datetime import date, datetime, timezone
from io import BytesIO
from unittest.mock import AsyncMock, patch

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.auth.models import Organization, User
from app.core.database import Base
from app.core.deps import get_db
from app.core.security import create_access_token, hash_password
from app.main import app

TEST_DATABASE_URL = "sqlite+aiosqlite:///:memory:"

ORG_A_ID = uuid.uuid4()
ORG_B_ID = uuid.uuid4()


# ---------------------------------------------------------------------------
# Fixtures (mirror test_reports_api.py — session-scoped engine)
# ---------------------------------------------------------------------------


@pytest_asyncio.fixture(scope="session")
async def engine():
    e = create_async_engine(TEST_DATABASE_URL, echo=False)
    async with e.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    yield e
    await e.dispose()


@pytest_asyncio.fixture
async def db_session(engine):
    async with engine.connect() as conn:
        await conn.begin()
        factory = async_sessionmaker(bind=conn, expire_on_commit=False, class_=AsyncSession)
        async with factory() as session:
            yield session
        await conn.rollback()


@pytest_asyncio.fixture
async def client(db_session):
    async def _override():
        yield db_session

    app.dependency_overrides[get_db] = _override
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c
    app.dependency_overrides.clear()


@pytest_asyncio.fixture
async def org_a(db_session):
    org = Organization(id=ORG_A_ID, name="Org A", slug="org-a-stk", plan="pro")
    db_session.add(org)
    await db_session.flush()
    return org


@pytest_asyncio.fixture
async def manager_user(db_session, org_a):
    user = User(
        id=uuid.uuid4(),
        org_id=org_a.id,
        email="manager-stk@org-a.com",
        password_hash=hash_password("pass"),
        role="manager",
    )
    db_session.add(user)
    await db_session.flush()
    return user


@pytest_asyncio.fixture
async def viewer_user(db_session, org_a):
    user = User(
        id=uuid.uuid4(),
        org_id=org_a.id,
        email="viewer-stk@org-a.com",
        password_hash=hash_password("pass"),
        role="viewer",
    )
    db_session.add(user)
    await db_session.flush()
    return user


def _auth_headers(user: User) -> dict:
    token = create_access_token(user.id, user.org_id, user.role)
    return {"Authorization": f"Bearer {token}"}


def _make_real_xlsx() -> bytes:
    from openpyxl import Workbook as WB

    wb = WB()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# Auth & parameter validation tests
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_stock_export_requires_authentication(client):
    resp = await client.get(
        "/api/v1/reports/stock-export",
        params={"date_from": "2026-01-01", "date_to": "2026-01-31"},
    )
    assert resp.status_code == 401


@pytest.mark.asyncio
async def test_stock_export_missing_dates_returns_422(client, manager_user):
    resp = await client.get(
        "/api/v1/reports/stock-export",
        headers=_auth_headers(manager_user),
    )
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_stock_export_date_to_before_date_from_returns_400(client, manager_user):
    resp = await client.get(
        "/api/v1/reports/stock-export",
        params={"date_from": "2026-03-01", "date_to": "2026-01-01"},
        headers=_auth_headers(manager_user),
    )
    assert resp.status_code == 400
    assert "INVALID_DATE_RANGE" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_stock_export_range_exceeds_limit_returns_400(client, manager_user):
    resp = await client.get(
        "/api/v1/reports/stock-export",
        params={"date_from": "2024-01-01", "date_to": "2026-01-01"},
        headers=_auth_headers(manager_user),
    )
    assert resp.status_code == 400
    assert "INVALID_DATE_RANGE" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_stock_export_returns_xlsx_content_type(client, manager_user):
    fake_bytes = _make_real_xlsx()
    with patch(
        "app.reports.router.service.build_stock_export",
        new_callable=AsyncMock,
        return_value=fake_bytes,
    ):
        resp = await client.get(
            "/api/v1/reports/stock-export",
            params={"date_from": "2026-01-01", "date_to": "2026-01-31"},
            headers=_auth_headers(manager_user),
        )
    assert resp.status_code == 200
    assert (
        resp.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "stock_2026-01-01_2026-01-31.xlsx" in resp.headers["content-disposition"]


@pytest.mark.asyncio
async def test_stock_export_viewer_role_allowed(client, viewer_user):
    fake_bytes = _make_real_xlsx()
    with patch(
        "app.reports.router.service.build_stock_export",
        new_callable=AsyncMock,
        return_value=fake_bytes,
    ):
        resp = await client.get(
            "/api/v1/reports/stock-export",
            params={"date_from": "2026-01-01", "date_to": "2026-01-31"},
            headers=_auth_headers(viewer_user),
        )
    assert resp.status_code == 200


# ---------------------------------------------------------------------------
# Unit tests: workbook structure
# ---------------------------------------------------------------------------


def _make_stock_rows(in_stock: bool | None = True) -> list:
    from app.reports.repository import StockRow

    return [
        StockRow(
            brand_name="TestBrand",
            sku_article="ART-001",
            sku_name="Test SKU",
            platform_name="WB",
            scored_at=date(2026, 1, 15),
            week_number=3,
            year=2026,
            in_stock=in_stock,
            warehouse_qty=50 if in_stock else 0,
            plan_tt_count=100,
        )
    ]


def test_stock_workbook_has_correct_headers():
    from app.reports.service import _build_stock_workbook

    wb = _build_stock_workbook([], date_from=date(2026, 1, 1), date_to=date(2026, 1, 31))
    ws = wb.active

    assert "2026-01-01" in (ws["A1"].value or "")

    headers = [ws.cell(row=2, column=c).value for c in range(1, 11)]
    assert "Бренд" in headers
    assert "В наличии" in headers
    assert "Остаток на складе" in headers
    assert "План (ТТ)" in headers


def test_stock_workbook_in_stock_row_is_green():
    from app.reports.service import _build_stock_workbook

    rows = _make_stock_rows(in_stock=True)
    wb = _build_stock_workbook(rows, date_from=date(2026, 1, 1), date_to=date(2026, 1, 31))
    ws = wb.active

    fill_rgb = ws.cell(row=3, column=1).fill.fgColor.rgb
    assert "C6EFCE" in fill_rgb


def test_stock_workbook_out_of_stock_row_is_red():
    from app.reports.service import _build_stock_workbook

    rows = _make_stock_rows(in_stock=False)
    wb = _build_stock_workbook(rows, date_from=date(2026, 1, 1), date_to=date(2026, 1, 31))
    ws = wb.active

    fill_rgb = ws.cell(row=3, column=1).fill.fgColor.rgb
    assert "FFC7CE" in fill_rgb


def test_stock_workbook_null_in_stock_no_fill():
    from app.reports.service import _build_stock_workbook

    rows = _make_stock_rows(in_stock=None)
    wb = _build_stock_workbook(rows, date_from=date(2026, 1, 1), date_to=date(2026, 1, 31))
    ws = wb.active

    assert ws.cell(row=3, column=1).fill.fill_type in (None, "none", "patternType")


def test_stock_workbook_missing_plan_shows_dash():
    from app.reports.repository import StockRow
    from app.reports.service import _build_stock_workbook

    rows = [
        StockRow(
            brand_name="B",
            sku_article="001",
            sku_name="N",
            platform_name="WB",
            scored_at=date(2026, 1, 15),
            week_number=3,
            year=2026,
            in_stock=True,
            warehouse_qty=10,
            plan_tt_count=None,
        )
    ]
    wb = _build_stock_workbook(rows, date_from=date(2026, 1, 1), date_to=date(2026, 1, 31))
    ws = wb.active

    # plan_tt_count is column 10
    assert ws.cell(row=3, column=10).value == "—"


def test_stock_workbook_in_stock_field_formatted_as_da_or_net():
    from app.reports.service import _build_stock_workbook

    rows_true = _make_stock_rows(in_stock=True)
    rows_false = _make_stock_rows(in_stock=False)

    wb_t = _build_stock_workbook(rows_true, date(2026, 1, 1), date(2026, 1, 31))
    wb_f = _build_stock_workbook(rows_false, date(2026, 1, 1), date(2026, 1, 31))

    # in_stock column is 8
    assert wb_t.active.cell(row=3, column=8).value == "Да"
    assert wb_f.active.cell(row=3, column=8).value == "Нет"
